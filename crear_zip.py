import zipfile
from pathlib import Path

# Contenido de los archivos del proyecto
requirements_content = """pymupdf==1.25.3
python-docx==1.1.2
"""

procesar_icfes_content = '''import os
import re
import json
import logging
from pathlib import Path
import fitz  # PyMuPDF
from docx import Document

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()]
)

BASE_DIR = Path(r"C:\\Users\\gonza\\OneDrive\\Desktop\\ICFES")
OUTPUT_DIR = BASE_DIR / "ICFES_Output"

def limpiar_texto(texto):
    if not texto:
        return ""
    return re.sub(r'\\s+', ' ', texto).strip()

def parsear_preguntas_pdf(pdf_path: Path):
    doc = fitz.open(pdf_path)
    BANCO_PAGINAS = len(doc)
    texto_completo_acumulado = ""
    
    for num_pag, pagina in enumerate(doc):
        txt = pagina.get_text("text")
        texto_completo_acumulado += f"\\n--- PAGINA {num_pag + 1} ---\\n" + txt

    es_escaneado = len(texto_completo_acumulado.strip()) < 200
    if es_escaneado:
        return [], BANCO_PAGINAS, True, ["Documento escaneado o sin texto seleccionable suficiente."]

    preguntas = []
    patron_pregunta = re.compile(r'(?:^|\\n)\\s*(\\d{1,3})\\s*[\\.\\-]\\s+', re.MULTILINE)
    coincidencias = list(patron_pregunta.finditer(texto_completo_acumulado))
    
    if not coincidencias:
        return [], BANCO_PAGINAS, True, ["No se pudieron detectar preguntas automáticamente mediante patrones estándar."]

    for i in range(len(coincidencias)):
        match_actual = coincidencias[i]
        num_preg = int(match_actual.group(1))
        inicio_bloque = match_actual.end()
        fin_bloque = coincidencias[i + 1].start() if i + 1 < len(coincidencias) else len(texto_completo_acumulado)
        
        contenido_bloque = texto_completo_acumulado[inicio_bloque:fin_bloque].strip()
        
        patron_opciones = re.compile(r'\\b([A-Da-d])[\\.\\)]\\s+(.*?)(?=\\b[A-Da-d][\\.\\)]\\s+|\\Z)', re.DOTALL)
        opciones_match = list(patron_opciones.finditer(contenido_bloque))
        
        opciones = []
        enunciado_y_contexto = contenido_bloque
        
        if opciones_match:
            primer_opt_pos = opciones_match[0].start()
            enunciado_y_contexto = contenido_bloque[:primer_opt_pos].strip()
            for opt_match in opciones_match:
                letra = opt_match.group(1).upper()
                texto_opcion = limpiar_texto(opt_match.group(2))
                opciones.append({"letra": letra, "texto": texto_opcion})
        
        requiere_revision = len(opciones) < 2
        obs_pregunta = ["Pregunta detectada con pocas opciones de respuesta."] if requiere_revision else []

        preguntas.append({
            "numero": num_preg,
            "enunciado": limpiar_texto(enunciado_y_contexto),
            "contexto": "",
            "opciones": opciones,
            "paginas_origen": [1],
            "elementos_visuales": [],
            "requiere_revision": requiere_revision,
            "observaciones": obs_pregunta
        })

    return preguntas, BANCO_PAGINAS, False, []

def generar_word(preguntas, metadata, output_docx_path: Path):
    doc = Document()
    titulo = f"ICFES - {metadata['anio']} | {metadata['convocatoria']} | {metadata['area']}"
    doc.add_heading(titulo, level=1)
    
    for p in preguntas:
        doc.add_paragraph(f"Pregunta {p['numero']}.", style='Heading 2')
        if p['enunciado']:
            doc.add_paragraph(p['enunciado'])
        for opt in p['opciones']:
            doc.add_paragraph(f"{opt['letra']}. {opt['texto']}", style='List Bullet')
        doc.add_paragraph()

    output_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_docx_path))

def procesar_coleccion(pdf_especifico: Path = None):
    if not BASE_DIR.exists():
        logging.error(f"La carpeta principal no existe: {BASE_DIR}")
        return

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    archivos_pdf = [pdf_especifico] if pdf_especifico else [p for p in BASE_DIR.glob("**/*.pdf") if "ICFES_Output" not in p.parts]

    logging.info(f"Se encontraron {len(archivos_pdf)} archivos PDF para procesar.")
    
    todas_las_preguntas_consolidado = []
    registro_control = []

    for pdf_path in archivos_pdf:
        logging.info(f"Procesando: {pdf_path.relative_to(BASE_DIR)}")
        
        try:
            rel_path = pdf_path.relative_to(BASE_DIR)
            partes = rel_path.parts
            anio = int(partes[0]) if len(partes) > 0 and partes[0].isdigit() else 0
            convocatoria = partes[1] if len(partes) > 1 else "Desconocida"
            area = pdf_path.stem.replace("_", " ").title()
        except Exception:
            anio, convocatoria, area = 0, "Desconocida", pdf_path.stem.replace("_", " ").title()

        preguntas, total_paginas, escaneado, observaciones = parsear_preguntas_pdf(pdf_path)
        
        resultado_json = {
            "archivo_origen": pdf_path.name,
            "anio": anio,
            "convocatoria": convocatoria,
            "area": area,
            "total_paginas": total_paginas,
            "total_preguntas_detectadas": len(preguntas),
            "preguntas": preguntas
        }
        
        out_sub_dir = OUTPUT_DIR / pdf_path.relative_to(BASE_DIR).parent
        out_sub_dir.mkdir(parents=True, exist_ok=True)
        
        with open(out_sub_dir / f"{pdf_path.stem}.json", "w", encoding="utf-8") as f:
            json.dump(resultado_json, f, ensure_ascii=False, indent=4)
            
        docx_path = out_sub_dir / f"{pdf_path.stem}.docx"
        generar_word(preguntas, {"anio": anio, "convocatoria": convocatoria, "area": area}, docx_path)
        
        for p in preguntas:
            p_cons = p.copy()
            p_cons["metadata_origen"] = {"anio": anio, "convocatoria": convocatoria, "area": area, "archivo_origen": pdf_path.name}
            todas_las_preguntas_consolidado.append(p_cons)
            
        registro_control.append({
            "archivo": pdf_path.name,
            "paginas": total_paginas,
            "preguntas_detectadas": len(preguntas),
            "escaneado": escaneado,
            "observaciones": observaciones
        })

    with open(OUTPUT_DIR / "todas_las_preguntas.json", "w", encoding="utf-8") as f:
        json.dump(todas_las_preguntas_consolidado, f, ensure_ascii=False, indent=4)
        
    with open(OUTPUT_DIR / "registro_control.json", "w", encoding="utf-8") as f:
        json.dump(registro_control, f, ensure_ascii=False, indent=4)
        
    logging.info("¡Procesamiento completado!")

if __name__ == "__main__":
    procesar_coleccion()
'''

# Generar archivos físicos y empaquetar
base_path = Path(r"C:\Users\gonza\OneDrive\Desktop\ICFES")
base_path.mkdir(parents=True, exist_ok=True)

req_file = base_path / "requirements.txt"
req_file.write_text(requirements_content, encoding="utf-8")

script_file = base_path / "procesar_icfes.py"
script_file.write_text(procesar_icfes_content, encoding="utf-8")

zip_path = base_path / "proyecto_icfes.zip"
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
    zipf.write(req_file, arcname="requirements.txt")
    zipf.write(script_file, arcname="procesar_icfes.py")

print(f"¡Archivo ZIP generado exitosamente en: {zip_path}")